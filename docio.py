"""
Document read / write.

extract_text(...)  -> plain text for the detection + review step.
redact_document(...) -> writes a de-identified file IN THE SAME FORMAT where
                        possible (.docx, .xlsx), preserving structure. PDFs are
                        extracted to a de-identified .docx because PDFs can't be
                        safely edited in place -- and a clean text/docx is what
                        you feed an AI anyway.
"""
from __future__ import annotations

import io
import os

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


def file_kind(filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    return {".docx": "docx", ".pdf": "pdf", ".xlsx": "xlsx", ".txt": "txt"}.get(ext, "")


# ---- extraction (for detection) ---------------------------------------------
def extract_text(data: bytes, kind: str) -> str:
    if kind == "txt":
        return data.decode("utf-8", errors="replace")
    if kind == "docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        for section in doc.sections:
            for hdrftr in (section.header, section.footer):
                parts.extend(p.text for p in hdrftr.paragraphs)
        return "\n".join(parts)
    if kind == "xlsx":
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        parts = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                parts.extend(str(c) for c in row if c is not None)
        return "\n".join(parts)
    if kind == "pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    raise ValueError(f"Unsupported kind: {kind}")


# ---- redacted output --------------------------------------------------------
def _redact_paragraph(paragraph, replace_fn) -> int:
    """Replace text in a paragraph. If anything changes we flatten the
    paragraph's runs into one (formatting within a changed line is lost, but the
    redaction is guaranteed correct -- the right trade-off for a privacy gate)."""
    original = paragraph.text
    if not original:
        return 0
    new, hits = replace_fn(original)
    if hits and new != original:
        for run in list(paragraph.runs):
            run.text = ""
        if paragraph.runs:
            paragraph.runs[0].text = new
        else:
            paragraph.add_run(new)
    return hits


def _redact_docx(data: bytes, replace_fn) -> tuple[bytes, int]:
    doc = Document(io.BytesIO(data))
    hits = 0
    for p in doc.paragraphs:
        hits += _redact_paragraph(p, replace_fn)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    hits += _redact_paragraph(p, replace_fn)
    for section in doc.sections:
        for hdrftr in (section.header, section.footer):
            for p in hdrftr.paragraphs:
                hits += _redact_paragraph(p, replace_fn)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), hits


def _redact_xlsx(data: bytes, replace_fn) -> tuple[bytes, int]:
    wb = load_workbook(io.BytesIO(data))  # keep formulas/formatting
    hits = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    new, n = replace_fn(cell.value)
                    if n:
                        cell.value = new
                        hits += n
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), hits


def _text_to_docx(text: str) -> bytes:
    doc = Document()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def redact_document(data: bytes, kind: str, replace_fn) -> tuple[bytes, str, int]:
    """Return (output_bytes, output_extension, hit_count)."""
    if kind == "docx":
        out, hits = _redact_docx(data, replace_fn)
        return out, ".docx", hits
    if kind == "xlsx":
        out, hits = _redact_xlsx(data, replace_fn)
        return out, ".xlsx", hits
    if kind == "txt":
        new, hits = replace_fn(data.decode("utf-8", errors="replace"))
        return new.encode("utf-8"), ".txt", hits
    if kind == "pdf":
        text = extract_text(data, "pdf")
        new, hits = replace_fn(text)
        return _text_to_docx(new), ".docx", hits
    raise ValueError(f"Unsupported kind: {kind}")
